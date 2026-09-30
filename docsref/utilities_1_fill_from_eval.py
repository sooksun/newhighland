# -*- coding: utf-8 -*-
"""
ขั้นที่ 1/2 — เติมข้อมูลไฟฟ้า/น้ำประปา/อินเทอร์เน็ต ลงชีต รายโรงเรียน จาก highland_eval / island_eval
ไฟล์: docs/ครู_DKSรวมนำเขาOBEC_Asset_รายโรงเรียน.xlsx (baseline 1,577 โรงเรียน)
กติกา: ใช้ข้อมูลปีล่าสุดที่มีคำตอบ (ถ้ามีทั้งสองตาราง ใช้ปีล่าสุด);
ไม่พบข้อมูลไฟฟ้า → ถือว่ามีไฟฟ้า (สมมติฐาน); ไม่พบข้อมูลน้ำ/อินเทอร์เน็ต → 'ไม่พบข้อมูล'
เขียนคอลัมน์ Q–Y (layout ชั่วคราว) → ต้องรัน utilities_2_province_summary.py ต่อทุกครั้ง
ต้องการ: pip install openpyxl pymysql ; MySQL ของ Laragon (root@localhost ไม่มีรหัส)
"""
import sys
sys.stdout.reconfigure(encoding='utf-8')

import os, shutil, re, collections
import openpyxl, pymysql
from copy import copy

SRC = r'D:\laragon\www\newhighland\docs\ครู_DKSรวมนำเขาOBEC_Asset_รายโรงเรียน.xlsx'
BAK = SRC.replace('.xlsx', '_สำรองก่อนเพิ่มไฟฟ้าประปา.xlsx')
if not os.path.exists(BAK):  # สำรองต้นฉบับครั้งแรกเท่านั้น — รันซ้ำจะไม่ทับไฟล์สำรอง
    shutil.copy2(SRC, BAK)

db = pymysql.connect(host='localhost', user='root', database='ssrainfo_ssra', charset='utf8mb4')
cur = db.cursor()

def clean(s):
    return re.sub(r'\s+', ' ', str(s).replace('<br>', ' ').replace('\\n', ' ')).strip()

cur.execute('SELECT id, master07 FROM citeria_master07'); H_WATER = {int(a): clean(b) for a, b in cur.fetchall()}
cur.execute('SELECT id, master08 FROM citeria_master08'); H_POWER = {int(a): clean(b) for a, b in cur.fetchall()}
cur.execute('SELECT id, master10 FROM citeria_master10'); H_NET = {int(a): clean(b) for a, b in cur.fetchall()}
# island ใช้ชุดตัวเลือกเดียวกัน (IslandOption '10','11','12')
I_POWER = {1: 'มีเฉพาะระบบไฟฟ้าจากแหล่งพลังงานทางเลือก (โซล่าเซลล์ ลม น้ำ เครื่องปั่นไฟ ฯลฯ)', 2: 'มีระบบไฟฟ้าส่วนภูมิภาค'}
I_WATER = {1: 'น้ำที่ต่อจากแหล่งน้ำธรรมชาติ เช่น ประปาภูเขา แม่น้ำ ลำธาร',
           2: 'น้ำที่สูบจาก บ่อน้ำ สระน้ำ บ่อน้ำบาดาล ของชุมชนหรือสาธารณะ',
           3: 'น้ำที่สูบจาก บ่อน้ำ สระน้ำ บ่อน้ำบาดาล ของโรงเรียน',
           4: 'น้ำประปา ที่จัดทำโดยชุมชน หรือ หมู่บ้าน',
           5: 'น้ำประปา ที่จัดทำโดยองค์กรปกครองส่วนท้องถิ่น',
           6: 'น้ำประปา ที่จัดทำโดยการประปาส่วนภูมิภาค'}
TAP = {4, 5, 6}  # รหัสที่เป็น "น้ำประปา"
I_NET = {1: 'ไม่มีเครือข่ายสัญญาณอินเทอร์เน็ต',
         2: 'มีเครือข่ายจากจานรับสัญญาณดาวเทียมเท่านั้น',
         3: 'มีเครือข่ายสัญญาณไร้สาย (โทรศัพท์มือถือ)',
         4: 'มีเครือข่ายสัญญาณแบบ Leased line หรือ Fiber Optic'}

def net_status(cs):
    """ใช้ระดับที่ดีที่สุดที่ตอบ: 4 Fiber/Leased line > 3 มือถือ > 2 ดาวเทียม > 1 ไม่มี"""
    if 4 in cs: return 'มี (Fiber/Leased line)'
    if 3 in cs: return 'มี (สัญญาณมือถือเท่านั้น)'
    if 2 in cs: return 'มี (เฉพาะดาวเทียม)'
    return 'ไม่มี'

def codes(v):
    return sorted({int(x) for x in re.findall(r'\d+', str(v or '')) if int(x) > 0})

# rec[sc_id] = {'power': (codes, src, year), 'water': (...), 'net': (...)}  เก็บเฉพาะปีล่าสุดที่มีคำตอบ ต่อหัวข้อ
rec = collections.defaultdict(dict)
def keep(sc, key, cs, src, year, labels):
    cs = [c for c in cs if c in labels]
    if not cs: return
    old = rec[sc].get(key)
    if old is None or year > old[2]:
        rec[sc][key] = (cs, src, year, labels)

cur.execute('SELECT sc_id, acadyears, citeria07, citeria08, citeria10 FROM highland_eval')
for sc, y, w, p, n in cur.fetchall():
    keep(str(sc), 'water', codes(w), 'พื้นที่สูง (highland_eval)', int(y), H_WATER)
    keep(str(sc), 'power', codes(p), 'พื้นที่สูง (highland_eval)', int(y), H_POWER)
    keep(str(sc), 'net', codes(n), 'พื้นที่สูง (highland_eval)', int(y), H_NET)
cur.execute('SELECT sc_id, acadyears, citeria11, citeria10, citeria12 FROM island_eval')
for sc, y, w, p, n in cur.fetchall():
    keep(str(sc), 'water', codes(w), 'พื้นที่เกาะ (island_eval)', int(y), I_WATER)
    keep(str(sc), 'power', codes(p), 'พื้นที่เกาะ (island_eval)', int(y), I_POWER)
    keep(str(sc), 'net', codes(n), 'พื้นที่เกาะ (island_eval)', int(y), I_NET)

wb = openpyxl.load_workbook(SRC)
ws = wb['รายโรงเรียน']
HEAD = ['ไฟฟ้า', 'ระบบไฟฟ้า (รายละเอียด)', 'แหล่งข้อมูลไฟฟ้า',
        'น้ำประปา', 'แหล่งน้ำ (รายละเอียด)', 'แหล่งข้อมูลน้ำ',
        'อินเทอร์เน็ต', 'ระบบอินเทอร์เน็ต (รายละเอียด)', 'แหล่งข้อมูลอินเทอร์เน็ต']
C0 = 17  # คอลัมน์ Q (ต่อจาก P=หมายเหตุ)
hdr = ws.cell(1, 16)
for k, h in enumerate(HEAD):
    c = ws.cell(1, C0 + k, h)
    c.font, c.fill, c.border, c.alignment = copy(hdr.font), copy(hdr.fill), copy(hdr.border), copy(hdr.alignment)

ASSUME = 'ไม่พบใน DB — ถือว่ามีไฟฟ้า (สมมติฐาน)'
NOT_FOUND = 'ไม่พบข้อมูลใน DB'
stat = collections.Counter()
for r in range(2, ws.max_row + 1):
    sc = str(ws.cell(r, 1).value or '').strip()
    if not sc.isdigit():  # ข้ามแถวว่าง/แถว 'รวม' ท้ายชีต
        for k in range(len(HEAD)): ws.cell(r, C0 + k).value = None
        continue
    d = rec.get(sc, {})
    p = d.get('power')
    if p:
        cs, src, y, lab = p
        pv = 'มี' if 2 in cs else 'มี (เฉพาะพลังงานทางเลือก)'
        vals = [pv, ' / '.join(lab[c] for c in cs), f'{src} ปี {y}']
        stat['power_db'] += 1; stat['power:' + pv] += 1
    else:
        vals = ['มี', ASSUME, 'สมมติฐาน']
        stat['power_assume'] += 1
    w = d.get('water')
    if w:
        cs, src, y, lab = w
        wv = 'มี' if TAP & set(cs) else 'ไม่มี'
        vals += [wv, ' / '.join(lab[c] for c in cs), f'{src} ปี {y}']
        stat['water_db'] += 1; stat['water:' + wv] += 1
    else:
        vals += [NOT_FOUND, '', '']
        stat['water_none'] += 1
    n = d.get('net')
    if n:
        cs, src, y, lab = n
        nv = net_status(cs)
        vals += [nv, ' / '.join(lab[c] for c in cs), f'{src} ปี {y}']
        stat['net_db'] += 1; stat['net:' + nv] += 1
    else:
        vals += [NOT_FOUND, '', '']
        stat['net_none'] += 1
    for k, v in enumerate(vals):
        ws.cell(r, C0 + k, v)

widths = [22, 55, 30, 14, 60, 30, 24, 55, 30]
for k, wdt in enumerate(widths):
    ws.column_dimensions[openpyxl.utils.get_column_letter(C0 + k)].width = wdt
last = openpyxl.utils.get_column_letter(C0 + len(HEAD) - 1)
ws.auto_filter.ref = f'A1:{last}{ws.max_row}'

wb.save(SRC)
print(dict(stat))
print('backup:', BAK)
