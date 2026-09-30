# -*- coding: utf-8 -*-
"""
ขั้นที่ 2/2 — รันต่อจาก utilities_1_fill_from_eval.py
1) จัดคอลัมน์ Q–X ใหม่: ไฟฟ้า | ระบบไฟฟ้า | น้ำประปา | แหล่งน้ำ | อินเทอร์เน็ต | ระบบอินเทอร์เน็ต
   | แหล่งข้อมูล | ปีที่สำรวจ (ตัวเลข)
2) สร้าง/แทนที่ชีต 'สรุปสาธารณูปโภครายจังหวัด' (สูตร COUNTIFS อ้างชีต รายโรงเรียน)
รันซ้ำได้: ถ้าคอลัมน์ถูกจัดแล้ว (X1 = 'ปีที่สำรวจ') จะข้ามขั้น 1 และสร้างชีตสรุปใหม่
"""
import sys
sys.stdout.reconfigure(encoding='utf-8')

import re, collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from copy import copy

F = r'D:\laragon\www\newhighland\docs\ครู_DKSรวมนำเขาOBEC_Asset_รายโรงเรียน.xlsx'
wb = openpyxl.load_workbook(F)
ws = wb['รายโรงเรียน']
LAST = 1578  # แถวสุดท้ายของโรงเรียน (1579 = รวม)
C0 = 17      # คอลัมน์ Q
NOSRC = 'สมมติฐาน (ไม่พบใน DB)'

# ---- 1) รายโรงเรียน: layout ชั่วคราวจากขั้น 1 (Q–Y, แหล่งข้อมูลแยกต่อหัวข้อ) → layout สุดท้าย (Q–X)
HEAD = ['ไฟฟ้า', 'ระบบไฟฟ้า (รายละเอียด)', 'น้ำประปา', 'แหล่งน้ำ (รายละเอียด)',
        'อินเทอร์เน็ต', 'ระบบอินเทอร์เน็ต (รายละเอียด)', 'แหล่งข้อมูล', 'ปีที่สำรวจ']
rows = []
DONE = ws.cell(1, C0 + 7).value == 'ปีที่สำรวจ'  # จัดคอลัมน์ไปแล้ว → อ่านค่าเดิม
if not DONE and ws.cell(1, C0 + 8).value != 'แหล่งข้อมูลอินเทอร์เน็ต':
    sys.exit('ไม่พบ layout จากขั้นที่ 1 — รัน utilities_1_fill_from_eval.py ก่อน')
for r in range(2, LAST + 1):
    if DONE:
        rows.append(tuple(ws.cell(r, c).value for c in range(C0, C0 + 8))); continue
    p, pd, ps, w, wd, wsrc, n, nd, ns = (ws.cell(r, c).value for c in range(C0, C0 + 9))
    # ทั้งสามหัวข้อมาจาก record เดียวกัน (ปีล่าสุด) — ถ้าต่างกันให้หยุด ไม่รวมแหล่งข้อมูลมั่ว
    srcs = {x for x in (ps, wsrc, ns) if x and x != 'สมมติฐาน'}
    assert len(srcs) <= 1, (r, srcs)
    m = re.match(r'(.*) ปี (\d{4})$', next(iter(srcs), ''))
    src, yr = (m.group(1), int(m.group(2))) if m else (NOSRC, None)
    rows.append((p, pd, w, wd, n, nd, src, yr))
for r, vals in enumerate(rows, start=2):
    for k, val in enumerate(vals):
        ws.cell(r, C0 + k).value = val  # ws.cell(..., None) ไม่ล้างค่าเดิม
    ws.cell(r, C0 + 8).value = None  # ล้างคอลัมน์ Y ที่เหลือจาก layout ชั่วคราว
    ws.cell(r, C0 + 7).alignment = Alignment(horizontal='center')
for k, h in enumerate(HEAD):
    ws.cell(1, C0 + k, h)
hdr = ws.cell(1, C0 + 8); hdr.value = None; hdr.fill = PatternFill(); hdr.border = Border()
for k, wdt in enumerate([24, 55, 18, 60, 26, 60, 28, 11]):
    ws.column_dimensions[L(C0 + k)].width = wdt
ws.auto_filter.ref = f'A1:{L(C0 + 7)}{ws.max_row}'

years = sorted({x[7] for x in rows if x[7]})

# ---- 2) ชีตสรุป
ref = wb['สรุปรายจังหวัด']
provs = [ref.cell(i, 1).value for i in range(2, ref.max_row + 1) if ref.cell(i, 1).value != 'รวม']
NAME = 'สรุปสาธารณูปโภครายจังหวัด'
for old in (NAME, 'สรุปไฟฟ้าประปารายจังหวัด'):  # ชื่อเดิมก่อนเพิ่มอินเทอร์เน็ต
    if old in wb.sheetnames:
        del wb[old]
sh = wb.create_sheet(NAME, index=1)

hf, hfill = copy(ref.cell(1, 1).font), copy(ref.cell(1, 1).fill)
thin = Side(style='thin', color='BFBFBF')
box = Border(left=thin, right=thin, top=thin, bottom=thin)
ctr = Alignment(horizontal='center', vertical='center', wrap_text=True)

R = lambda col: f"รายโรงเรียน!${col}$2:${col}${LAST}"
PV, Q, S, U, W, X = R('C'), R('Q'), R('S'), R('U'), R('W'), R('X')
cnt = lambda rng, crit: (lambda i: f'=COUNTIFS({PV},$A{i},{rng},"{crit}")')

# (หัวกลุ่ม, หัวคอลัมน์, สูตรแถว i) — ตำแหน่งคอลัมน์ A..Q คงที่ (สูตรร้อยละอ้างตัวอักษรตรง ๆ)
G_P, G_W, G_N, G_Y = 'ไฟฟ้า', 'น้ำประปา', 'อินเทอร์เน็ต', 'ปีที่สำรวจ (จำนวนโรงเรียน)'
cols = [
    (None, 'จังหวัด', None),                                                   # A
    (None, 'โรงเรียนทั้งหมด', lambda i: f'=COUNTIF({PV},$A{i})'),                # B
    (None, 'พบข้อมูลใน DB', lambda i: f'=COUNTIFS({PV},$A{i},{X},">0")'),        # C
    (G_P, 'มีไฟฟ้าส่วนภูมิภาค', lambda i: f'=COUNTIFS({PV},$A{i},{Q},"มี")-F{i}'),  # D (มี ลบ สมมติฐาน)
    (G_P, 'เฉพาะพลังงานทางเลือก', cnt(Q, 'มี (เฉพาะพลังงานทางเลือก)')),        # E
    (G_P, 'ถือว่ามีไฟฟ้า (สมมติฐาน)', cnt(R('R'), '*สมมติฐาน*')),               # F
    (G_P, 'รวมมีไฟฟ้า', lambda i: f'=SUM(D{i}:F{i})'),                          # G
    (G_W, 'มีน้ำประปา', cnt(S, 'มี')),                                          # H
    (G_W, 'ไม่มีน้ำประปา', cnt(S, 'ไม่มี')),                                     # I
    (G_W, 'ไม่พบข้อมูล', cnt(S, 'ไม่พบข้อมูล*')),                                # J
    (G_W, 'ร้อยละมีประปา (ของที่พบข้อมูล)', lambda i: f'=IF(H{i}+I{i}=0,"-",H{i}/(H{i}+I{i}))'),  # K
    (G_N, 'Fiber / Leased line', cnt(U, 'มี (Fiber/Leased line)')),            # L
    (G_N, 'สัญญาณมือถือเท่านั้น', cnt(U, 'มี (สัญญาณมือถือเท่านั้น)')),          # M
    (G_N, 'ดาวเทียมเท่านั้น', cnt(U, 'มี (เฉพาะดาวเทียม)')),                     # N
    (G_N, 'ไม่มีอินเทอร์เน็ต', cnt(U, 'ไม่มี')),                                 # O
    (G_N, 'ไม่พบข้อมูล', cnt(U, 'ไม่พบข้อมูล*')),                                # P
    (G_N, 'ร้อยละมีอินเทอร์เน็ต (ของที่พบข้อมูล)',
     lambda i: f'=IF(SUM(L{i}:O{i})=0,"-",SUM(L{i}:N{i})/SUM(L{i}:O{i}))'),     # Q
] + [(G_Y, f'ปี {y}', (lambda y: lambda i: f'=COUNTIFS({PV},$A{i},{X},{y})')(y)) for y in years] + [
    (G_Y, 'ปีล่าสุด', lambda i: f'=IF(C{i}=0,"-",SUMPRODUCT(MAX(({PV}=$A{i})*{X})))'),
]
NC = len(cols)
PCT = {11, 17}  # K, Q

# หัวตาราง 2 แถว: กลุ่ม (merge) + ชื่อคอลัมน์
c = 1
while c <= NC:
    g = cols[c - 1][0]
    if g is None:
        sh.cell(1, c, cols[c - 1][1]); sh.merge_cells(start_row=1, start_column=c, end_row=2, end_column=c); c += 1
    else:
        e = c
        while e < NC and cols[e][0] == g: e += 1
        sh.cell(1, c, g); sh.merge_cells(start_row=1, start_column=c, end_row=1, end_column=e)
        for k in range(c, e + 1): sh.cell(2, k, cols[k - 1][1])
        c = e + 1
for rr in (1, 2):
    for k in range(1, NC + 1):
        cell = sh.cell(rr, k); cell.font, cell.fill, cell.alignment, cell.border = copy(hf), copy(hfill), ctr, box
sh.row_dimensions[2].height = 45

for n, p in enumerate(provs):
    i = 3 + n
    sh.cell(i, 1, p).border = box
    for k in range(2, NC + 1):
        cell = sh.cell(i, k, cols[k - 1][2](i)); cell.border = box; cell.alignment = Alignment(horizontal='center')
        if k in PCT: cell.number_format = '0.0%'
last = 2 + len(provs); tot = last + 1
sh.cell(tot, 1, 'รวม')
for k in range(2, NC + 1):
    col = L(k)
    if k in PCT:
        f = cols[k - 1][2](tot)  # ร้อยละคิดจากยอดรวม ไม่ใช่รวมร้อยละ
    elif k == NC:
        f = f'=MAX({X})'
    else:
        f = f'=SUM({col}3:{col}{last})'
    cell = sh.cell(tot, k, f); cell.alignment = Alignment(horizontal='center')
    if k in PCT: cell.number_format = '0.0%'
for k in range(1, NC + 1):
    cell = sh.cell(tot, k); cell.font = Font(bold=True); cell.border = box
    cell.fill = PatternFill('solid', fgColor='DDEBF7')

notes = [
    'หมายเหตุ',
    '• ข้อมูลจากฐานข้อมูลคัดกรองโรงเรียนพื้นที่สูง (highland_eval ข้อ 7 น้ำ / ข้อ 8 ไฟฟ้า / ข้อ 10 อินเทอร์เน็ต) '
    'และพื้นที่เกาะ (island_eval ข้อ 4.1 ไฟฟ้า / 4.2 น้ำ / 4.3 อินเทอร์เน็ต) ใช้ปีล่าสุดที่มีคำตอบ',
    '• โรงเรียนที่ไม่พบข้อมูลใน DB ถือว่ามีไฟฟ้า (สมมติฐาน) ส่วนน้ำประปาและอินเทอร์เน็ตระบุเป็น "ไม่พบข้อมูล"',
    '• "มีน้ำประปา" = ประปาหมู่บ้าน/ชุมชน, ประปา อปท., ประปาส่วนภูมิภาค ; ประปาภูเขา/แหล่งน้ำธรรมชาติ/บ่อ/บาดาล นับเป็น "ไม่มีน้ำประปา"',
    '• อินเทอร์เน็ตนับตามระดับดีที่สุดที่ตอบ: Fiber/Leased line > สัญญาณมือถือ > ดาวเทียม > ไม่มี',
    '• ร้อยละคำนวณจากโรงเรียนที่พบข้อมูลเท่านั้น ; ตัวเลขคำนวณจากชีต รายโรงเรียน (คอลัมน์ Q–X) อัตโนมัติ',
]
for k, t in enumerate(notes):
    cell = sh.cell(tot + 2 + k, 1, t)
    if k == 0: cell.font = Font(bold=True)

sh.column_dimensions['A'].width = 18
for k in range(2, NC + 1):
    sh.column_dimensions[L(k)].width = 13
sh.freeze_panes = 'B3'
wb.calculation.fullCalcOnLoad = True
wb.save(F)

# ---- ตรวจค่าที่ควรได้ (คำนวณใน Python) เทียบกับสูตร
cnt_ = collections.Counter()
for (p, pd, w, _, n, _, src, yr) in rows:
    cnt_['n'] += 1; cnt_['db'] += bool(yr)
    cnt_['assume'] += 'สมมติฐาน' in (pd or '')
    cnt_['pea'] += p == 'มี' and 'สมมติฐาน' not in (pd or ''); cnt_['alt'] += p == 'มี (เฉพาะพลังงานทางเลือก)'
    cnt_['tap'] += w == 'มี'; cnt_['notap'] += w == 'ไม่มี'; cnt_['w_nf'] += w.startswith('ไม่พบ')
    cnt_['net:' + n] += 1
    if yr: cnt_[yr] += 1
print(dict(cnt_)); print('years', years, 'provinces', len(provs), 'sheet cols', NC)
